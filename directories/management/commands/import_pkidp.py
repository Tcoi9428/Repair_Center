from pathlib import Path
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from directories.forms import CompanyForm
from directories.models import Company
from directories.services import record_change

FIELDS = ['code', 'company_name_full', 'company_name_top_full', 'company_name_short', 'company_name_top_short', 'location_explotation_short', 'company_location', 'company_contact', 'company_ceo', 'service_contract', 'service_contract_num']

def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

class Command(BaseCommand):
    help = 'Импортирует предприятия из Excel. Сначала --dry-run. Существующие коды пропускаются без перезаписи.'

    def add_arguments(self, parser):
        parser.add_argument('path')
        parser.add_argument('--dry-run', action='store_true')

    @transaction.atomic
    def handle(self, *args, **options):
        try:
            workbook = load_workbook(Path(options['path']), read_only=True, data_only=True)
        except Exception as exc:
            raise CommandError(f'Не удалось прочитать Excel: {exc}') from exc
        created = skipped = 0
        seen = set()
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = [cell_text(value) for value in next(rows)]
            if len(set(headers)) != len(headers) or set(headers) != set(FIELDS):
                raise CommandError('Столбцы Excel должны точно соответствовать 11 полям ПКиДП.')
            for number, row in enumerate(rows, 2):
                if not any(value is not None and str(value).strip() for value in row):
                    continue
                data = dict(zip(headers, map(cell_text, row)))
                code = data['code']
                if code in seen:
                    raise CommandError(f'Строка {number}: повтор кода {code}. Импорт отменен.')
                seen.add(code)
                flag = data['service_contract'].casefold()
                if flag not in {'да', 'нет', 'true', 'false', '1', '0'}:
                    raise CommandError(f'Строка {number}: в service_contract ожидается Да или Нет.')
                data['service_contract'] = 'True' if flag in {'да', 'true', '1'} else 'False'
                if Company.objects.filter(code=code).exists():
                    skipped += 1
                    continue
                form = CompanyForm(data)
                if not form.is_valid():
                    raise CommandError(f'Строка {number}: {form.errors.as_text()}. Импорт отменен.')
                company = form.save()
                record_change(company, None, {}, 'create')
                created += 1
        finally:
            workbook.close()
        if options['dry_run']:
            transaction.set_rollback(True)
        self.stdout.write(self.style.SUCCESS(f'{"Проверка без сохранения" if options["dry_run"] else "Импорт завершен"}: новых {created}, пропущено существующих {skipped}.'))

