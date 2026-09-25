import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class ParsedOperation:
    sequence: int
    section: str
    operation_number: str
    description: str
    standard_minutes: int


@dataclass(frozen=True)
class ParsedMaterial:
    sequence: int
    kind: str
    nomenclature_number: str
    name: str
    quantity: Decimal | None
    unit: str
    notes: str


@dataclass(frozen=True)
class ParsedTechnologyCard:
    sheet_name: str
    interval_hours: int
    operations: tuple[ParsedOperation, ...]
    materials: tuple[ParsedMaterial, ...]

    @property
    def total_minutes(self):
        return sum(item.standard_minutes for item in self.operations)


def clean_text(value):
    if value is None:
        return ''
    return ' '.join(str(value).replace('\xa0', ' ').split())


def number_text(value):
    if isinstance(value, bool) or value is None:
        return ''
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = clean_text(value)
    return text if text.isdecimal() else ''


def parse_quantity(value, default_unit):
    if value is None or value == '':
        return None, default_unit
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return Decimal(str(value)), default_unit
    text = clean_text(value).replace(',', '.')
    match = re.match(r'^([0-9]+(?:\.[0-9]+)?)\s*(.*)$', text)
    if not match:
        return None, default_unit
    try:
        quantity = Decimal(match.group(1))
    except InvalidOperation:
        return None, default_unit
    return quantity, clean_text(match.group(2)) or default_unit


def find_row(sheet, phrase, start=1):
    phrase = phrase.casefold()
    for row in range(start, sheet.max_row + 1):
        for column in range(1, sheet.max_column + 1):
            if phrase in clean_text(sheet.cell(row, column).value).casefold():
                return row
    raise ValueError(f'На листе «{sheet.title}» не найден блок «{phrase}».')


def parse_sheet(sheet):
    interval_match = re.search(r'(\d+)', sheet.title)
    if not interval_match:
        raise ValueError(f'В названии листа «{sheet.title}» не найден интервал ТО.')
    interval_hours = int(interval_match.group(1))
    operations_header = find_row(sheet, 'Перечень работ')
    materials_title = find_row(sheet, 'Список запасных частей', operations_header + 1)
    lubricants_title = find_row(sheet, 'Перечень ГСМ', materials_title + 1)

    operations = []
    section = ''
    for row in range(operations_header + 1, materials_title):
        operation_number = number_text(sheet.cell(row, 1).value)
        description = clean_text(sheet.cell(row, 2).value)
        if operation_number and description:
            raw_minutes = sheet.cell(row, 10).value
            try:
                minutes = int(Decimal(str(raw_minutes).replace(',', '.')))
            except (InvalidOperation, TypeError, ValueError):
                raise ValueError(
                    f'Лист «{sheet.title}», строка {row}: не указано корректное время операции.'
                )
            operations.append(
                ParsedOperation(
                    sequence=len(operations) + 1,
                    section=section,
                    operation_number=operation_number,
                    description=description,
                    standard_minutes=minutes,
                )
            )
        elif clean_text(sheet.cell(row, 1).value) and not description:
            candidate = clean_text(sheet.cell(row, 1).value)
            if candidate not in {'№ п/п'}:
                section = candidate

    materials = []
    part_sequence = 0
    for row in range(materials_title + 2, lubricants_title):
        if not number_text(sheet.cell(row, 1).value):
            continue
        name = clean_text(sheet.cell(row, 4).value)
        if not name:
            continue
        part_sequence += 1
        quantity, unit = parse_quantity(sheet.cell(row, 10).value, 'шт')
        materials.append(
            ParsedMaterial(
                sequence=part_sequence,
                kind='part',
                nomenclature_number=clean_text(sheet.cell(row, 2).value),
                name=name,
                quantity=quantity,
                unit=unit,
                notes='',
            )
        )

    lubricant_sequence = 0
    for row in range(lubricants_title + 2, sheet.max_row + 1):
        if not number_text(sheet.cell(row, 1).value):
            continue
        name = clean_text(sheet.cell(row, 2).value)
        if not name:
            continue
        lubricant_sequence += 1
        quantity, unit = parse_quantity(sheet.cell(row, 10).value, 'л')
        materials.append(
            ParsedMaterial(
                sequence=lubricant_sequence,
                kind='lubricant',
                nomenclature_number='',
                name=name,
                quantity=quantity,
                unit=unit,
                notes=clean_text(sheet.cell(row, 4).value),
            )
        )

    if not operations:
        raise ValueError(f'На листе «{sheet.title}» операции не найдены.')
    return ParsedTechnologyCard(
        sheet_name=sheet.title,
        interval_hours=interval_hours,
        operations=tuple(operations),
        materials=tuple(materials),
    )


def parse_workbook(path):
    path = Path(path)
    workbook = load_workbook(path, data_only=True, read_only=False)
    return tuple(parse_sheet(sheet) for sheet in workbook.worksheets if sheet.sheet_state == 'visible')


def build_markdown_summary(source, cards):
    lines = [
        '# Извлечение технологических карт ПДМ 10 ШААЗ',
        '',
        f'Источник: `{Path(source).name}`.',
        '',
        'MarkItDown отсутствовал в окружении. Книга прочитана резервным способом через `openpyxl`.',
        '',
        '| Лист | Вид ТО | Операций | Норматив, мин | Материалов |',
        '|---|---:|---:|---:|---:|',
    ]
    for card in cards:
        lines.append(
            f'| {card.sheet_name} | ТО-{card.interval_hours} | {len(card.operations)} | '
            f'{card.total_minutes} | {len(card.materials)} |'
        )
    lines.extend(['', '## Правила извлечения', ''])
    lines.extend(
        [
            '- Заголовок операций определяется по тексту «Перечень работ».',
            '- Строки с номером в колонке A, описанием в B и временем в J считаются операциями.',
            '- Строки без номера между операциями считаются названиями разделов.',
            '- Запасные части извлекаются из отдельной таблицы с каталожным номером, наименованием и количеством.',
            '- ГСМ извлекаются отдельно; место применения сохраняется в примечании.',
            '- Пустые номенклатурные номера не заполняются вымышленными значениями.',
        ]
    )
    return '\n'.join(lines) + '\n'
